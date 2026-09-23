#!/usr/bin/env python3
"""Read-only CoreMap/MIMU/Myanmar Post reconciliation audit.

This script never connects to PostgreSQL. Production evidence must first be
exported by phase1_export_readonly.sql inside a READ ONLY transaction.
"""

from __future__ import annotations

import argparse
import collections
import csv
import difflib
import hashlib
import io
import json
import re
import subprocess
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pandas as pd

csv.field_size_limit(100_000_000)


MIMU_VERSION = "9.7 Jan 2026"
MIMU_GEOMETRY_VERSION = "9.4"
MIMU_SOURCE_URL = "https://themimu.info/place-codes"
POSTAL_VERSION = "V-1.0 Sep 2021"
POSTAL_SOURCE_URL = "https://github.com/MyanmarPost/MyanmarPostalCode"
OSM_SOURCE = "tools/data-pipeline/local-osm/data/osm/myanmar-260823.osm.pbf"

SR_PARENT_REMAP = {
    "MMR007": "MMR111",
    "MMR008": "MMR111",
    "MMR014": "MMR222",
    "MMR015": "MMR222",
    "MMR016": "MMR222",
}

MYANMAR_RE = re.compile(r"[\u1000-\u109f\uaa60-\uaa7f\ua9e0-\ua9ff]")
PCODE_RE = re.compile(r"^MMR[0-9A-Z]+$")


@dataclass
class SourceRow:
    level: str
    pcode: str
    parent_pcode: str
    name_en: str
    name_my: str
    subtype: str
    gad_status: str
    mapping_status: str
    official: bool
    source_sheet: str


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value).strip()


def compact_name(value: Any, level: str = "") -> str:
    text = unicodedata.normalize("NFKC", clean(value)).casefold()
    text = text.replace("ဧရာဝတီ", "ဧရာဝတီ").replace("nay pyi taw", "naypyitaw")
    text = text.replace("ayeyawady", "ayeyarwady").replace("irrawaddy", "ayeyarwady")
    suffixes = [
        r"\bunion territory\b", r"\bstate\b", r"\bregion\b", r"\bdistrict\b",
        r"\btownship\b", r"\btown\b", r"\bquarter\b", r"\bward\b",
        r"\bvillage[ -]?tract\b", r"\bself[ -]?administered (?:zone|division)\b",
        "ပြည်ထောင်စုနယ်မြေ", "တိုင်းဒေသကြီး", "ပြည်နယ်", "ခရိုင်", "မြို့နယ်",
        "ကျေးရွာအုပ်စု", "ရပ်ကွက်", "မြို့",
    ]
    for suffix in suffixes:
        text = re.sub(suffix, " ", text)
    text = re.sub(r"\b(?:no|number)\s*[.(]?\s*([0-9]+)\s*[).]?", r"\1", text)
    text = text.replace("အမှတ်", "")
    text = re.sub(r"[^0-9a-z\u1000-\u109f\uaa60-\uaa7f\ua9e0-\ua9ff]+", "", text)
    return text


def contains_myanmar(value: str) -> bool:
    return bool(MYANMAR_RE.search(value or ""))


def bool_value(value: Any) -> bool:
    return clean(value).lower() in {"t", "true", "1", "yes"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sheet_dicts(workbook: openpyxl.Workbook, sheet: str, header_row: int) -> list[dict[str, Any]]:
    ws = workbook[sheet]
    headers = [clean(cell.value) for cell in ws[header_row]]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        row = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        if any(clean(value) for value in values):
            rows.append(row)
    return rows


def first(row: dict[str, Any], *keys: str) -> str:
    for key in keys:
        if clean(row.get(key)):
            return clean(row.get(key))
    return ""


def read_mimu(workbook_path: Path) -> tuple[list[SourceRow], list[dict[str, str]], dict[str, str]]:
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True, keep_vba=True)
    sources: list[SourceRow] = []

    for row in sheet_dicts(wb, "01_SR", 1):
        status = first(row, "GAD_State/Region_Status")
        sources.append(SourceRow(
            "state_region", first(row, "SR_Pcode"), "", first(row, "SR_Name_Eng"),
            first(row, "SR_Name_MMR"), "", status,
            first(row, "MIMU_State/Region_Mapping_Status"), status == "Active", "01_SR",
        ))

    for row in sheet_dicts(wb, "SAD_SAZ", 1):
        status = first(row, "GAD_SAD/SAZ_Status")
        name_en = first(row, "SAD/SAZ_Name_Eng")
        subtype = "self_administered_division" if "Division" in name_en else "self_administered_zone"
        sources.append(SourceRow(
            "self_administered_zone", first(row, "SAD/SAZ_Pcode"),
            SR_PARENT_REMAP.get(first(row, "SR_Pcode"), first(row, "SR_Pcode")),
            name_en, first(row, "SAD/SAZ_Name_MMR"), subtype, status,
            first(row, "MIMU_SAD/SAZ_Mapping_Status"), status == "Active", "SAD_SAZ",
        ))

    for row in sheet_dicts(wb, "02_District", 1):
        status = first(row, "GAD_District_Status")
        sources.append(SourceRow(
            "district", first(row, "District_Pcode"),
            SR_PARENT_REMAP.get(first(row, "SR"), first(row, "SR")),
            first(row, "District_Name_Eng"), first(row, "District_Name_MMR"), "district",
            status, first(row, "MIMU_District_Mapping_Status"), status == "Active", "02_District",
        ))

    for row in sheet_dicts(wb, "03_Township", 6):
        status = first(row, "GAD_Township_Status")
        sources.append(SourceRow(
            "township", first(row, "Tsp_Pcode"), first(row, "District/SAZ_Pcode"),
            first(row, "Township_Name_Eng"), first(row, "Township_Name_MMR"), "township",
            status, first(row, "MIMU_Township_Mapping_Status"), status == "Active", "03_Township",
        ))

    for sheet, subtype, pcode_key, en_key, my_key, gad_key, map_key in [
        ("_05_Ward", "ward", "Ward_Pcode", "Ward_Name_Eng", "Ward_Name_MMR", "GAD_VillageTract_Status", "MIMU_War_Mapping_Status"),
        ("_06-VillageTract", "village_tract", "VT_Pcode", "Village_Tract_Name_Eng", "Village_Tract_Name_MMR", "GAD_VillageTract_Status", "MIMU_VillageTract_Mapping_Status"),
    ]:
        for row in sheet_dicts(wb, sheet, 6):
            status = first(row, gad_key)
            sources.append(SourceRow(
                "ward_village_tract", first(row, pcode_key), first(row, "Tsp_Pcode"),
                first(row, en_key), first(row, my_key), subtype, status, first(row, map_key),
                status == "Active", sheet,
            ))

    towns: list[dict[str, str]] = []
    for row in sheet_dicts(wb, "04_Town", 6):
        if first(row, "GAD_Town_Status") == "Active":
            towns.append({
                "town_pcode": first(row, "Town_Pcode"),
                "township_pcode": first(row, "Tsp_Pcode"),
                "state_pcode": SR_PARENT_REMAP.get(first(row, "SR_Pcode"), first(row, "SR_Pcode")),
                "name_en": first(row, "Town_Name_Eng"),
                "name_my": first(row, "Town_Name_MMR"),
            })

    metadata = {}
    for row in wb["Metadata"].iter_rows(values_only=True):
        if clean(row[0]).endswith(":"):
            metadata[clean(row[0]).rstrip(":")] = clean(row[1])
    return sources, towns, metadata


def run_ogr_geojson(shapefile: Path) -> list[dict[str, Any]]:
    result = subprocess.run(
        ["ogr2ogr", "-f", "GeoJSONSeq", "/vsistdout/", str(shapefile)],
        check=True, capture_output=True, text=True,
    )
    return [json.loads(line.lstrip("\x1e")) for line in result.stdout.splitlines() if line.strip()]


def prepare_geometry_csv(mimu_root: Path, output_csv: Path) -> None:
    specs = [
        ("state_region", "mmr_polbnda_adm1_250k_mimu_1/mmr_polbnda_adm1_250k_mimu_1.shp", "ST_PCODE", "", "ST", "ST_MMR"),
        ("district", "mmr_polbnda_adm2_250k_mimu/mmr_polbnda_adm2_250k_mimu.shp", "DT_PCODE", "ST_PCODE", "DT", "DT_MMR"),
        ("township", "mmr_polbnda_adm3_250k_mimu_1/mmr_polbnda_adm3_250k_mimu_1.shp", "TS_PCODE", "DT_PCODE", "TS", "TS_MMR"),
    ]
    rows = []
    for level, relative, pcode_key, parent_key, en_key, my_key in specs:
        for feature in run_ogr_geojson(mimu_root / "unzipped" / relative):
            props = feature.get("properties") or {}
            parent = clean(props.get(parent_key)) if parent_key else ""
            if level == "district":
                parent = SR_PARENT_REMAP.get(parent, parent)
            rows.append({
                "source_level": level,
                "pcode": clean(props.get(pcode_key)),
                "parent_pcode": parent,
                "name_en": clean(props.get(en_key)),
                "name_my": clean(props.get(my_key)),
                "source_version": MIMU_GEOMETRY_VERSION,
                "geom_geojson": json.dumps(feature.get("geometry"), ensure_ascii=False, separators=(",", ":")),
            })
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def prepare_spatial_sql(geometry_csv: Path, output_sql: Path, candidates_csv: Path) -> None:
    """Generate a SELECT-only PostGIS audit query using an inline VALUES CTE."""
    rows = list(csv.DictReader(geometry_csv.open(encoding="utf-8")))

    def quoted(value: str, tag: str) -> str:
        if f"${tag}$" in value:
            raise ValueError(f"Unexpected dollar-quote delimiter in {tag}")
        return f"${tag}${value}${tag}$"

    values = []
    for index, row in enumerate(rows):
        tag = f"g{index}"
        values.append(
            "(" + ",".join([
                quoted(row["source_level"], f"l{index}"),
                quoted(row["pcode"], f"p{index}"),
                quoted(row["parent_pcode"], f"r{index}"),
                quoted(row["name_en"], f"e{index}"),
                quoted(row["name_my"], f"m{index}"),
                quoted(row["source_version"], f"v{index}"),
                quoted(row["geom_geojson"], tag),
            ]) + ")"
        )
    sql = f"""\\set ON_ERROR_STOP on
\\pset footer off
\\pset format csv
BEGIN TRANSACTION READ ONLY;
\\o {candidates_csv}
WITH input(source_level,pcode,parent_pcode,name_en,name_my,source_version,geom_geojson) AS (
    VALUES
    {','.join(values)}
),
source_geometry AS (
    SELECT source_level,pcode,parent_pcode,name_en,name_my,source_version,
           ST_Multi(ST_CollectionExtract(ST_MakeValid(ST_SetSRID(ST_GeomFromGeoJSON(geom_geojson),4326)),3)) AS geom
    FROM input
),
candidate_base AS (
    SELECT s.source_level,s.pcode,s.parent_pcode,s.name_en,s.name_my,s.source_version,
           s.geom AS source_geom,a.geom AS core_geom,
           a.id AS core_id,a.parent_id AS core_parent_id,a.canonical_name AS core_name,
           md5(ST_AsEWKB(a.geom)) AS core_geom_md5,
           ST_Covers(a.geom,ST_PointOnSurface(s.geom)) AS source_point_in_core,
           ST_Covers(s.geom,ST_PointOnSurface(a.geom)) AS core_point_in_source,
           round(ST_Distance(ST_PointOnSurface(s.geom)::geography,ST_PointOnSurface(a.geom)::geography)::numeric,2) AS point_distance_m
    FROM source_geometry s
    JOIN ref.ref_admin_levels l ON l.code=s.source_level
    JOIN core.core_admin_areas a ON a.admin_level_id=l.id AND a.is_active AND a.deleted_at IS NULL
                              AND a.geom && s.geom
),
ranked_base AS (
    SELECT c.*,row_number() OVER (PARTITION BY source_level,pcode ORDER BY
        (source_point_in_core AND core_point_in_source) DESC,
        source_point_in_core DESC,core_point_in_source DESC,point_distance_m,core_id) AS base_rank
    FROM candidate_base c
)
SELECT source_level,pcode,parent_pcode,name_en,name_my,source_version,
       core_id,core_parent_id,core_name,core_geom_md5,source_point_in_core,core_point_in_source,
       NULL::numeric AS source_covered_by_core_pct,NULL::numeric AS core_covered_by_source_pct,
       point_distance_m,base_rank AS spatial_rank
FROM ranked_base WHERE base_rank<=5 ORDER BY source_level,pcode,spatial_rank;
ROLLBACK;
"""
    output_sql.write_text(sql, encoding="utf-8")


def recursively_find_pcodes(value: Any) -> set[str]:
    found: set[str] = set()
    if isinstance(value, dict):
        for key, item in value.items():
            if "pcode" in str(key).casefold() and PCODE_RE.match(clean(item)):
                found.add(clean(item))
            found.update(recursively_find_pcodes(item))
    elif isinstance(value, list):
        for item in value:
            found.update(recursively_find_pcodes(item))
    return found


def read_osm_candidates(path: Path) -> dict[str, list[dict[str, Any]]]:
    candidates: dict[str, list[dict[str, Any]]] = collections.defaultdict(list)
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            feature = json.loads(line.lstrip("\x1e"))
            props = feature.get("properties") or {}
            pcode = clean(props.get("pcode"))
            geometry_type = clean((feature.get("geometry") or {}).get("type"))
            if pcode and geometry_type in {"Polygon", "MultiPolygon"}:
                candidates[pcode].append({
                    "feature_id": clean(feature.get("id")),
                    "name": first(props, "name:en", "name", "name:my"),
                    "name_en": first(props, "name:en"),
                    "name_my": first(props, "name:my"),
                    "admin_level": first(props, "admin_level"),
                    "geometry_type": geometry_type,
                    "source": first(props, "source"),
                })
    return candidates


def source_expected_type(source: SourceRow) -> str:
    if source.level == "state_region":
        if source.pcode == "MMR018":
            return "union_territory"
        if source.pcode in {"MMR001", "MMR002", "MMR003", "MMR004", "MMR011", "MMR012", "MMR222"}:
            return "state"
        return "region"
    return source.subtype or source.level


def names_for_area(names: pd.DataFrame) -> dict[int, list[dict[str, Any]]]:
    grouped: dict[int, list[dict[str, Any]]] = collections.defaultdict(list)
    for record in names.to_dict("records"):
        record["admin_area_id"] = int(record["admin_area_id"])
        grouped[record["admin_area_id"]].append(record)
    return grouped


def area_all_names(area: dict[str, Any], grouped_names: dict[int, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    values = list(grouped_names.get(int(area["id"]), []))
    values.append({
        "name": clean(area.get("canonical_name")), "language_code": "und",
        "name_type": "canonical", "is_primary": False,
    })
    return values


def primary_name(area_id: int, language: str, grouped_names: dict[int, list[dict[str, Any]]]) -> str:
    rows = [row for row in grouped_names.get(area_id, []) if clean(row.get("language_code")) == language and bool_value(row.get("is_primary"))]
    rows.sort(key=lambda row: (-int(float(clean(row.get("search_weight")) or 0)), int(row["id"])))
    return clean(rows[0]["name"]) if rows else ""


def spatial_map(path: Path | None) -> dict[tuple[str, str], list[dict[str, Any]]]:
    if not path or not path.exists():
        return {}
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    result: dict[tuple[str, str], list[dict[str, Any]]] = collections.defaultdict(list)
    for row in frame.to_dict("records"):
        result[(row["source_level"], row["pcode"])].append(row)
    return result


def float_or_zero(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def overlap_text(row: dict[str, Any] | None) -> str:
    if not row:
        return "not_available"
    return (
        f"source_in_core={row.get('source_covered_by_core_pct','')}%;"
        f"core_in_source={row.get('core_covered_by_source_pct','')}%;"
        f"source_point_in_core={row.get('source_point_in_core','')};"
        f"core_point_in_source={row.get('core_point_in_source','')};"
        f"distance_m={row.get('point_distance_m','')}"
    )


def strong_spatial(row: dict[str, Any]) -> bool:
    source_pct = float_or_zero(row.get("source_covered_by_core_pct"))
    core_pct = float_or_zero(row.get("core_covered_by_source_pct"))
    mutual = bool_value(row.get("source_point_in_core")) and bool_value(row.get("core_point_in_source"))
    distance = float_or_zero(row.get("point_distance_m"))
    return (mutual and distance <= 50_000) or (min(source_pct, core_pct) >= 55 and max(source_pct, core_pct) >= 95)


def match_admin(
    sources: list[SourceRow], areas: pd.DataFrame, names: pd.DataFrame,
    spatial: dict[tuple[str, str], list[dict[str, Any]]], osm: dict[str, list[dict[str, Any]]],
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    area_records = areas.to_dict("records")
    for area in area_records:
        area["id"] = int(area["id"])
        area["parent_id"] = int(area["parent_id"]) if clean(area.get("parent_id")) else None
    area_by_id = {area["id"]: area for area in area_records}
    by_level: dict[str, list[dict[str, Any]]] = collections.defaultdict(list)
    for area in area_records:
        by_level[clean(area["admin_level"])].append(area)
    grouped_names = names_for_area(names)
    country_id = next(area["id"] for area in by_level["country"] if bool_value(area["is_active"]) and not clean(area["deleted_at"]))

    source_ref_index: dict[str, list[int]] = collections.defaultdict(list)
    for area in area_records:
        try:
            refs = json.loads(clean(area.get("source_refs")) or "{}")
        except json.JSONDecodeError:
            refs = {}
        for pcode in recursively_find_pcodes(refs):
            source_ref_index[pcode].append(area["id"])

    pcode_to_core: dict[str, int] = {}
    matched_core: set[int] = set()
    reports: list[dict[str, Any]] = []
    source_order = {"state_region": 1, "district": 2, "self_administered_zone": 2, "township": 3, "ward_village_tract": 4}

    for source in sorted(sources, key=lambda row: (source_order[row.level], row.pcode)):
        parent_core = country_id if source.level == "state_region" else pcode_to_core.get(source.parent_pcode)
        # Parent identity is a hard gate.  Searching every same-level row when a
        # parent is unresolved creates dangerous false positives for repeated
        # names such as "Ward 1" and "Village Tract".
        active_candidates = [
            area for area in by_level[source.level]
            if bool_value(area.get("is_active")) and not clean(area.get("deleted_at"))
            and (source.level == "state_region" or (parent_core is not None and area.get("parent_id") == parent_core))
        ]
        candidate_ids: list[int] = []
        method = ""
        reason = ""

        pcode_ids = [area_id for area_id in source_ref_index.get(source.pcode, []) if area_by_id[area_id]["admin_level"] == source.level]
        if len(pcode_ids) == 1:
            candidate_ids, method = pcode_ids, "source_refs_pcode"
        elif len(pcode_ids) > 1:
            candidate_ids, method = pcode_ids, "duplicate_source_refs_pcode"

        def exact_candidates(target: str, language: str, alias_only: bool = False) -> list[int]:
            norm_target = compact_name(target, source.level)
            if not norm_target:
                return []
            found = []
            for area in active_candidates:
                for name_row in area_all_names(area, grouped_names):
                    name_type = clean(name_row.get("name_type")).casefold()
                    lang = clean(name_row.get("language_code"))
                    if alias_only and name_type not in {"alternate", "alt", "alias", "short", "local"}:
                        continue
                    if not alias_only and language and lang not in {language, "und", ""}:
                        continue
                    if compact_name(name_row.get("name"), source.level) == norm_target:
                        found.append(area["id"])
                        break
            return sorted(set(found))

        if not candidate_ids:
            candidate_ids = exact_candidates(source.name_my, "my")
            method = "exact_normalized_my_parent" if candidate_ids else ""
        if not candidate_ids:
            candidate_ids = exact_candidates(source.name_en, "en")
            method = "exact_normalized_en_parent" if candidate_ids else ""
        if not candidate_ids:
            aliases = sorted(set(exact_candidates(source.name_my, "", True) + exact_candidates(source.name_en, "", True)))
            if aliases:
                candidate_ids, method = aliases, "existing_alias_exact_parent"

        spatial_rows = spatial.get((source.level, source.pcode), [])
        top_spatial = spatial_rows[0] if spatial_rows else None
        if len(candidate_ids) > 1 and top_spatial and strong_spatial(top_spatial):
            spatial_id = int(top_spatial["core_id"])
            if spatial_id in candidate_ids:
                candidate_ids, method = [spatial_id], method + "+strong_spatial_tiebreak"
        if not candidate_ids and top_spatial and strong_spatial(top_spatial):
            candidate_ids, method = [int(top_spatial["core_id"])], "strong_spatial"

        selected: dict[str, Any] | None = area_by_id[candidate_ids[0]] if len(candidate_ids) == 1 else None
        action = ""
        confidence = 0
        geometry_evidence = overlap_text(next((row for row in spatial_rows if selected and int(row["core_id"]) == selected["id"]), None))
        proposed_geometry_source = "preserve_core"

        if not source.official:
            action = "reference_only"
            confidence = 100 if selected else 80
            reason = f"MIMU GAD status={source.gad_status or 'blank'}; excluded from official target"
        elif len(candidate_ids) > 1:
            action = "duplicate_merge_candidate"
            confidence = 40
            reason = "Multiple exact CoreMap candidates under the same matched parent: " + ",".join(map(str, candidate_ids))
        elif selected:
            matched_core.add(selected["id"])
            pcode_to_core[source.pcode] = selected["id"]
            expected_type = source_expected_type(source)
            if selected.get("parent_id") != parent_core and parent_core is not None:
                action, confidence, reason = "update_parent", 95, "Exact identity match but parent differs from canonical MIMU hierarchy"
            elif clean(selected.get("admin_area_type")) not in {expected_type, "", "unknown"} or not bool_value(selected.get("is_official_boundary")):
                action, confidence, reason = "reclassify", 95, f"Core type/official flags differ from expected {expected_type}"
            else:
                primary_en = primary_name(selected["id"], "en", grouped_names)
                primary_my = primary_name(selected["id"], "my", grouped_names)
                if primary_en != source.name_en or primary_my != source.name_my:
                    action, confidence, reason = "update_names", 95, "Identity and parent match; MIMU primary English/Myanmar names differ"
                else:
                    action, confidence, reason = "exact_keep", 100, "Identity, parent, official classification and primary bilingual names already match"
        else:
            fuzzy: list[tuple[float, int, str]] = []
            targets = [compact_name(source.name_en, source.level), compact_name(source.name_my, source.level)]
            for area in active_candidates:
                for name_row in area_all_names(area, grouped_names):
                    candidate_name = compact_name(name_row.get("name"), source.level)
                    if candidate_name:
                        score = max(difflib.SequenceMatcher(None, target, candidate_name).ratio() for target in targets if target)
                        if score >= 0.80:
                            fuzzy.append((score, area["id"], clean(name_row.get("name"))))
            fuzzy.sort(reverse=True)
            osm_polygons = osm.get(source.pcode, [])
            if fuzzy:
                action, confidence = "manual_review", round(fuzzy[0][0] * 100)
                reason = "Fuzzy-name candidates are review-only: " + "; ".join(f"{item[1]}:{item[2]}:{item[0]:.3f}" for item in fuzzy[:3])
                method = "fuzzy_review_only"
            elif osm_polygons:
                action, confidence = "create_missing", 100
                method = "osm_pcode_exact_polygon"
                proposed_geometry_source = "osm_odbl"
                reason = "No CoreMap match; exact MIMU PCode exists on OSM polygon " + ",".join(item["feature_id"] for item in osm_polygons)
                geometry_evidence = "exact_pcode_osm_polygon;" + ";".join(f"{item['feature_id']}:{item['geometry_type']}" for item in osm_polygons)
            else:
                action, confidence = "blocked_missing_legal_geometry", 100
                method = "no_legal_geometry_found"
                proposed_geometry_source = "none"
                reason = "No CoreMap or exact-PCode OSM polygon; MIMU geometry is unavailable at this level or restricted to comparison/QA"
                if source.level == "ward_village_tract":
                    geometry_evidence = "MIMU v9.7 workbook has names/PCodes/hierarchy but no ward/village-tract polygon geometry"

        if selected and source.pcode not in pcode_to_core:
            pcode_to_core[source.pcode] = selected["id"]
        current_parent = area_by_id.get(selected.get("parent_id")) if selected and selected.get("parent_id") else None
        target_parent = area_by_id.get(parent_core) if parent_core else None
        reports.append({
            "source_level": source.level,
            "source_subtype": source.subtype,
            "source_official": source.official,
            "source_gad_status": source.gad_status,
            "mimu_pcode": source.pcode,
            "parent_pcode": source.parent_pcode,
            "coremap_id": selected["id"] if selected else "",
            "candidate_coremap_ids": ";".join(map(str, candidate_ids)),
            "current_name": clean(selected.get("canonical_name")) if selected else "",
            "target_name_en": source.name_en,
            "target_name_my": source.name_my,
            "current_parent_id": selected.get("parent_id") if selected else "",
            "current_parent_name": clean(current_parent.get("canonical_name")) if current_parent else "",
            "target_parent_id": parent_core or "",
            "target_parent_name": clean(target_parent.get("canonical_name")) if target_parent else "",
            "match_method": method,
            "geometry_overlap_evidence": geometry_evidence,
            "proposed_geometry_source": proposed_geometry_source,
            "proposed_action": action,
            "confidence_score": confidence,
            "review_reason": reason,
            "core_geom_md5_before": clean(selected.get("geom_md5")) if selected else "",
        })

    # Core-only rows that the source explicitly requires handling.
    for area_id in [5985, 5986, 6675, 6734, 6735]:
        area = area_by_id.get(area_id)
        if area:
            reports.append({
                "source_level": clean(area["admin_level"]), "source_subtype": clean(area.get("admin_area_type")),
                "source_official": False, "source_gad_status": "not_in_mimu_official_population",
                "mimu_pcode": "", "parent_pcode": "", "coremap_id": area_id, "candidate_coremap_ids": str(area_id),
                "current_name": clean(area["canonical_name"]), "target_name_en": "", "target_name_my": "",
                "current_parent_id": area.get("parent_id") or "", "current_parent_name": "", "target_parent_id": "",
                "target_parent_name": "", "match_method": "confirmed_foreign_core_id", "geometry_overlap_evidence": "outside Myanmar",
                "proposed_geometry_source": "preserve_for_audit_then_disable", "proposed_action": "disable_foreign",
                "confidence_score": 100, "review_reason": "Confirmed foreign polygon; inspect/reassign or exclude all dependencies before disabling",
                "core_geom_md5_before": clean(area.get("geom_md5")),
            })

    unmatched_active_state_rows = [
        area for area in by_level["state_region"]
        if bool_value(area.get("is_active")) and not clean(area.get("deleted_at")) and area["id"] not in matched_core
    ]
    for area in unmatched_active_state_rows:
        if area["id"] in {5985, 5986, 6675, 6734, 6735}:
            continue
        reports.append({
            "source_level": "state_region", "source_subtype": clean(area.get("admin_area_type")), "source_official": False,
            "source_gad_status": "core_only", "mimu_pcode": "", "parent_pcode": "", "coremap_id": area["id"],
            "candidate_coremap_ids": str(area["id"]), "current_name": clean(area["canonical_name"]), "target_name_en": "",
            "target_name_my": "", "current_parent_id": area.get("parent_id") or "", "current_parent_name": "",
            "target_parent_id": "", "target_parent_name": "", "match_method": "core_only_first_level",
            "geometry_overlap_evidence": "not_applicable", "proposed_geometry_source": "preserve_core",
            "proposed_action": "reclassify", "confidence_score": 100,
            "review_reason": "Core-only Wa/de-facto first-level area must become special_area/reference_only and leave official counts",
            "core_geom_md5_before": clean(area.get("geom_md5")),
        })
    return reports, pcode_to_core


def duplicate_report(areas: pd.DataFrame, names: pd.DataFrame) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    names2 = names.copy()
    names2["normalized_name"] = names2["name"].map(compact_name)
    grouped = names2.groupby(["admin_area_id", "language_code", "normalized_name"], dropna=False)
    for key, group in grouped:
        if len(group) > 1:
            rows.append({
                "duplicate_type": "exact_duplicate_name_rows", "admin_level": "", "parent_id": "",
                "normalized_name": key[2], "admin_area_ids": str(key[0]),
                "name_row_ids": ";".join(group["id"].astype(str)), "row_count": len(group),
                "review_reason": "Remove exact duplicate name rows while preserving one row and useful aliases",
            })
    area_level = {str(row["id"]): clean(row["admin_level"]) for row in areas.to_dict("records")}
    area_parent = {str(row["id"]): clean(row["parent_id"]) for row in areas.to_dict("records")}
    memberships: dict[tuple[str, str, str], set[str]] = collections.defaultdict(set)
    for row in names2.to_dict("records"):
        area_id = str(row["admin_area_id"])
        memberships[(area_level.get(area_id, ""), area_parent.get(area_id, ""), row["normalized_name"])].add(area_id)
    for (level, parent_id, normalized), ids in memberships.items():
        if normalized and len(ids) > 1:
            rows.append({
                "duplicate_type": "admin_area_merge_candidate", "admin_level": level, "parent_id": parent_id,
                "normalized_name": normalized, "admin_area_ids": ";".join(sorted(ids, key=int)), "name_row_ids": "",
                "row_count": len(ids), "review_reason": "Same normalized name, level and parent; geometry and dependencies require review before merge",
            })
    return rows


def read_postal(zip_path: Path) -> tuple[dict[str, list[dict[str, str]]], dict[str, list[dict[str, str]]], dict[str, Any]]:
    with zipfile.ZipFile(zip_path) as archive:
        names = {Path(name).name: name for name in archive.namelist()}
        frames: dict[str, list[dict[str, str]]] = {}
        hashes = {}
        for language, filename in [("en", "Myanmar_Locations_Postal_Code_EN.csv"), ("my", "Myanmar_Locations_Postal_Code_MM.csv")]:
            raw = archive.read(names[filename])
            hashes[language] = hashlib.sha256(raw).hexdigest()
            text = raw.decode("utf-8-sig")
            frames[language] = [{key: clean(value) for key, value in row.items()} for row in csv.DictReader(io.StringIO(text))]
    grouped = {}
    for language, rows in frames.items():
        by_code: dict[str, list[dict[str, str]]] = collections.defaultdict(list)
        for row in rows:
            by_code[row["Postal Code"]].append(row)
        grouped[language] = by_code
    all_codes = set(grouped["en"]) | set(grouped["my"])
    valid = {code for code in all_codes if re.fullmatch(r"\d{7}", code)}
    stats = {
        "source_rows_en": len(frames["en"]), "source_rows_my": len(frames["my"]),
        "valid_rows_en": sum(len(grouped["en"][code]) for code in valid),
        "valid_rows_my": sum(len(grouped["my"][code]) for code in valid),
        "unique_valid_codes": len(valid),
        "duplicate_extra_rows_en": sum(max(0, len(grouped["en"][code]) - 1) for code in valid),
        "duplicate_extra_rows_my": sum(max(0, len(grouped["my"][code]) - 1) for code in valid),
        "malformed_codes": sorted(code for code in all_codes if not re.fullmatch(r"\d{7}", code)),
        "sha256_en_csv": hashes["en"], "sha256_my_csv": hashes["my"], "sha256_zip": sha256_file(zip_path),
    }
    return grouped["en"], grouped["my"], stats


def join_values(rows: Iterable[dict[str, str]], key: str) -> str:
    return " | ".join(sorted({clean(row.get(key)) for row in rows if clean(row.get(key))}))


def postal_audit(
    en_by_code: dict[str, list[dict[str, str]]], my_by_code: dict[str, list[dict[str, str]]],
    sources: list[SourceRow], towns: list[dict[str, str]], pcode_to_core: dict[str, int],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    official_townships = [row for row in sources if row.level == "township" and row.official]
    official_locals = [row for row in sources if row.level == "ward_village_tract" and row.official]
    state_by_child: dict[str, str] = {}
    source_by_pcode = {row.pcode: row for row in sources}
    for township in official_townships:
        parent = source_by_pcode.get(township.parent_pcode)
        state_by_child[township.pcode] = parent.parent_pcode if parent and parent.level in {"district", "self_administered_zone"} else ""
    state_names: dict[str, str] = {}
    for state in [row for row in sources if row.level == "state_region" and row.official]:
        state_names[compact_name(state.name_en, "state_region")] = state.pcode
        state_names[compact_name(state.name_my, "state_region")] = state.pcode

    township_index: dict[tuple[str, str], set[str]] = collections.defaultdict(set)
    for township in official_townships:
        state = state_by_child.get(township.pcode, "")
        township_index[(state, compact_name(township.name_en, "township"))].add(township.pcode)
        township_index[(state, compact_name(township.name_my, "township"))].add(township.pcode)
    town_index: dict[tuple[str, str], set[str]] = collections.defaultdict(set)
    for town in towns:
        town_index[(town["state_pcode"], compact_name(town["name_en"], "town"))].add(town["township_pcode"])
        town_index[(town["state_pcode"], compact_name(town["name_my"], "town"))].add(town["township_pcode"])
    locals_by_township: dict[str, list[SourceRow]] = collections.defaultdict(list)
    for local in official_locals:
        locals_by_township[local.parent_pcode].append(local)

    valid_codes = sorted(code for code in set(en_by_code) | set(my_by_code) if re.fullmatch(r"\d{7}", code))
    reports: list[dict[str, Any]] = []
    manual: list[dict[str, Any]] = []
    for code in valid_codes:
        en_rows, my_rows = en_by_code.get(code, []), my_by_code.get(code, [])
        region_en, region_my = join_values(en_rows, "Region"), join_values(my_rows, "Region")
        township_en, township_my = join_values(en_rows, "Town / Township"), join_values(my_rows, "Town / Township")
        locality_en, locality_my = join_values(en_rows, "Quarter / Village Tract"), join_values(my_rows, "Quarter / Village Tract")
        state_candidates = {
            state_names.get(compact_name(value, "state_region"), "")
            for value in [region_en, region_my] if value
        } - {""}
        state_pcode = next(iter(state_candidates)) if len(state_candidates) == 1 else ""
        township_candidates: set[str] = set()
        for value in [township_en, township_my]:
            if value:
                township_candidates.update(township_index.get((state_pcode, compact_name(value, "township")), set()))
        method = "exact_township_name"
        if not township_candidates:
            for value in [township_en, township_my]:
                if value:
                    township_candidates.update(town_index.get((state_pcode, compact_name(value, "town")), set()))
            method = "exact_town_to_parent_township" if township_candidates else ""
        township_pcode = next(iter(township_candidates)) if len(township_candidates) == 1 else ""
        township_core_id = pcode_to_core.get(township_pcode, "")

        inferred_type = ""
        joined_locality = f"{locality_en} {locality_my}".casefold()
        if "quarter" in joined_locality or "ward" in joined_locality or "ရပ်ကွက်" in joined_locality:
            inferred_type = "ward"
        elif "village tract" in joined_locality or "village-tract" in joined_locality or "ကျေးရွာအုပ်စု" in joined_locality:
            inferred_type = "village_tract"
        local_candidates: list[SourceRow] = []
        if township_pcode:
            target_norms = {compact_name(locality_en, "ward_village_tract"), compact_name(locality_my, "ward_village_tract")} - {""}
            local_candidates = [
                local for local in locals_by_township[township_pcode]
                if target_norms & {compact_name(local.name_en, local.level), compact_name(local.name_my, local.level)}
            ]
        if inferred_type:
            typed = [local for local in local_candidates if local.subtype == inferred_type]
            if typed:
                local_candidates = typed
        local_pcode = local_candidates[0].pcode if len(local_candidates) == 1 else ""
        local_core_id = pcode_to_core.get(local_pcode, "")
        type_conflict = bool(len(local_candidates) == 1 and inferred_type and local_candidates[0].subtype != inferred_type)

        if len(state_candidates) != 1 or len(township_candidates) != 1:
            status = "ambiguous" if state_candidates or township_candidates else "unmatched_postal_locality"
        elif not township_core_id:
            status = "township_source_match_core_unresolved"
        elif len(local_candidates) > 1:
            status = "ambiguous_local_admin"
        elif type_conflict:
            status = "type_conflict"
        elif local_core_id:
            status = "exact_local_admin_match"
        else:
            status = "township_only"
        row = {
            "postal_code": code,
            "source_row_count_en": len(en_rows), "source_row_count_my": len(my_rows),
            "region_name_en": region_en, "region_name_my": region_my,
            "township_name_en": township_en, "township_name_my": township_my,
            "locality_name_en": locality_en, "locality_name_my": locality_my,
            "locality_type": inferred_type,
            "township_pcode": township_pcode, "township_admin_area_id": township_core_id,
            "local_pcode": local_pcode, "local_admin_area_id": local_core_id,
            "match_status": status, "match_method": method,
            "candidate_township_pcodes": ";".join(sorted(township_candidates)),
            "candidate_local_pcodes": ";".join(sorted(local.pcode for local in local_candidates)),
            "review_reason": "" if status in {"exact_local_admin_match", "township_only"} else "Identity is not proven; do not guess or create geometry",
        }
        reports.append(row)
        if status not in {"exact_local_admin_match", "township_only"}:
            manual.append(row)

    for malformed in sorted((set(en_by_code) | set(my_by_code)) - set(valid_codes)):
        manual.append({
            "postal_code": malformed, "source_row_count_en": len(en_by_code.get(malformed, [])),
            "source_row_count_my": len(my_by_code.get(malformed, [])), "region_name_en": join_values(en_by_code.get(malformed, []), "Region"),
            "region_name_my": join_values(my_by_code.get(malformed, []), "Region"),
            "township_name_en": join_values(en_by_code.get(malformed, []), "Town / Township"),
            "township_name_my": join_values(my_by_code.get(malformed, []), "Town / Township"),
            "locality_name_en": join_values(en_by_code.get(malformed, []), "Quarter / Village Tract"),
            "locality_name_my": join_values(my_by_code.get(malformed, []), "Quarter / Village Tract"),
            "locality_type": "", "township_pcode": "", "township_admin_area_id": "", "local_pcode": "",
            "local_admin_area_id": "", "match_status": "malformed_postal_code", "match_method": "seven_digit_validation",
            "candidate_township_pcodes": "", "candidate_local_pcodes": "", "review_reason": "Malformed code quarantined; no correction guessed",
        })
    return reports, manual


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = list(rows[0]) if rows else []
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        if fieldnames:
            writer.writeheader()
            writer.writerows(rows)


def summary_markdown(
    admin_rows: list[dict[str, Any]], postal_rows: list[dict[str, Any]], postal_stats: dict[str, Any],
    baseline: pd.DataFrame, foreign_dependencies: pd.DataFrame, fks: pd.DataFrame,
) -> str:
    official = [row for row in admin_rows if row["source_official"] is True and row["mimu_pcode"]]
    counts = collections.Counter((row["source_level"], row["proposed_action"]) for row in official)
    levels = ["state_region", "district", "self_administered_zone", "township", "ward_village_tract"]
    actions = sorted({action for _, action in counts})
    lines = [
        "# Administrative and postal reconciliation audit", "",
        "Read-only audit date: 2026-09-21. No production rows, schemas, migrations, tiles, or search indexes were changed.", "",
        "## Source populations and proposed actions", "",
        "| Level | Official MIMU rows | " + " | ".join(actions) + " | Required five-action subtotal | Gap |",
        "|---|---:|" + "---:|" * len(actions) + "---:|---:|",
    ]
    required_actions = {"exact_keep", "update_names", "update_parent", "reclassify", "create_missing"}
    for level in levels:
        total = sum(value for (item_level, _), value in counts.items() if item_level == level)
        subtotal = sum(counts[(level, action)] for action in required_actions)
        values = [str(counts[(level, action)]) for action in actions]
        lines.append(f"| {level} | {total} | " + " | ".join(values) + f" | {subtotal} | {total - subtotal} |")
    lines.extend([
        "", "The requested five-action equation is shown explicitly. Any gap is composed of `manual_review`, `duplicate_merge_candidate`, or `blocked_missing_legal_geometry`; treating those rows as approved creates would be unsafe.",
        "", "Township official target: **330**. The source workbook contains exactly 330 rows with `GAD_Township_Status=Active`.",
        "", "## Rechecked production baseline", "",
        "| Metric | Value | Detail |", "|---|---:|---|",
    ])
    for row in baseline.to_dict("records"):
        lines.append(f"| {row['metric']} | {row['value']} | {clean(row.get('detail'))} |")
    lines.extend(["", "## Postal audit", ""])
    for key, value in postal_stats.items():
        lines.append(f"- {key}: {value}")
    postal_counts = collections.Counter(row["match_status"] for row in postal_rows)
    exact_township = sum(
        1 for row in postal_rows
        if row["township_admin_area_id"] and row["match_method"] == "exact_township_name"
    )
    exact_town_parent = sum(
        1 for row in postal_rows
        if row["township_admin_area_id"] and row["match_method"] == "exact_town_to_parent_township"
    )
    exact_local = postal_counts["exact_local_admin_match"]
    ambiguous = postal_counts["ambiguous"] + postal_counts["ambiguous_local_admin"]
    unmatched = postal_counts["unmatched_postal_locality"] + postal_counts["township_source_match_core_unresolved"]
    lines.extend([
        "", "### Required postal reconciliation rollup", "",
        "| Metric | Count |", "|---|---:|",
        f"| Exact township-name matches to CoreMap | {exact_township} |",
        f"| Exact town-name to parent-township matches | {exact_town_parent} |",
        f"| Exact WVT matches to CoreMap | {exact_local} |",
        "| Alias matches | 0 |",
        f"| Ambiguous (township or local) | {ambiguous} |",
        f"| Unmatched / Core township unresolved | {unmatched} |",
        f"| Type conflicts | {postal_counts['type_conflict']} |",
    ])
    lines.extend(["", "### Postal match status", "", "| Status | Count |", "|---|---:|"])
    for status, count in sorted(postal_counts.items()):
        lines.append(f"| {status} | {count} |")
    region_counts = collections.Counter(row["region_name_en"] or row["region_name_my"] for row in postal_rows)
    lines.extend(["", "### Valid codes by source region", "", "| Region | Count |", "|---|---:|"])
    for region, count in sorted(region_counts.items()):
        lines.append(f"| {region} | {count} |")
    township_counts = collections.Counter(
        (row["region_name_en"] or row["region_name_my"], row["township_name_en"] or row["township_name_my"])
        for row in postal_rows
    )
    lines.extend(["", "### Valid codes by source region and township", "", "| Region | Township / town | Count |", "|---|---|---:|"])
    for (region, township), count in sorted(township_counts.items()):
        lines.append(f"| {clean(region).replace('|','/')} | {clean(township).replace('|','/')} | {count} |")
    lines.extend(["", "## Foreign polygon dependencies (must be resolved before disable)", ""])
    foreign_dependencies = foreign_dependencies[
        pd.to_numeric(foreign_dependencies.get("admin_area_id"), errors="coerce").notna()
    ] if not foreign_dependencies.empty else foreign_dependencies
    if foreign_dependencies.empty:
        lines.append("No dependencies found in the audited runtime tables.")
    else:
        lines.extend(["| Admin ID | Source | Column | Rows |", "|---:|---|---|---:|"])
        for row in foreign_dependencies.to_dict("records"):
            lines.append(f"| {row['admin_area_id']} | {row['source_table']} | {row['source_column']} | {row['row_count']} |")
    lines.extend(["", f"Foreign-key inventory: {len(fks)} constraints reference `core.core_admin_areas`; see the audit source evidence and do not merge/disable until every dependent table is handled."])

    lines.extend([
        "", "## Examples by non-empty action", "",
        "Ten examples are shown for each category that contains at least ten rows; where a category has fewer than ten rows, every available row is shown.", "",
    ])
    grouped: dict[str, list[dict[str, Any]]] = collections.defaultdict(list)
    for row in admin_rows:
        grouped[row["proposed_action"]].append(row)
    for action in sorted(grouped):
        lines.extend([f"### {action}", "", "| Level | PCode | Core ID | Current | Target EN | Reason |", "|---|---|---:|---|---|---|"])
        for row in grouped[action][:10]:
            reason = clean(row["review_reason"]).replace("|", "/")
            lines.append(f"| {row['source_level']} | {row['mimu_pcode']} | {row['coremap_id']} | {clean(row['current_name']).replace('|','/')} | {clean(row['target_name_en']).replace('|','/')} | {reason} |")
        lines.append("")
    lines.extend([
        "## Stop gate", "",
        "Phase 2 has not run. Review the action counts, duplicate decisions, fuzzy/manual candidates, blocked geometry rows, and foreign dependency plan before approving any migration or production data change.",
    ])
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    prepare = subparsers.add_parser("prepare-geometry")
    prepare.add_argument("--mimu-root", type=Path, required=True)
    prepare.add_argument("--output", type=Path, required=True)
    spatial_sql = subparsers.add_parser("prepare-spatial-sql")
    spatial_sql.add_argument("--geometry-csv", type=Path, required=True)
    spatial_sql.add_argument("--output", type=Path, required=True)
    spatial_sql.add_argument("--candidates-output", type=Path, required=True)

    report = subparsers.add_parser("report")
    report.add_argument("--mimu-root", type=Path, required=True)
    report.add_argument("--postal-zip", type=Path, required=True)
    report.add_argument("--core-dir", type=Path, required=True)
    report.add_argument("--osm-admin", type=Path, required=True)
    report.add_argument("--spatial-candidates", type=Path)
    report.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    if args.command == "prepare-geometry":
        prepare_geometry_csv(args.mimu_root, args.output)
        return
    if args.command == "prepare-spatial-sql":
        prepare_spatial_sql(args.geometry_csv, args.output, args.candidates_output)
        return

    workbook_path = args.mimu_root / "pcodes" / "Myanmar_PCodes_Release_9.7_Jan2026_StRgn_Dist_Tsp_Town_Ward_VT.xlsm"
    sources, towns, _metadata = read_mimu(workbook_path)
    areas = pd.read_csv(args.core_dir / "core_areas.csv", dtype=str, keep_default_na=False)
    names = pd.read_csv(args.core_dir / "core_names.csv", dtype=str, keep_default_na=False)
    baseline = pd.read_csv(args.core_dir / "baseline.csv", dtype=str, keep_default_na=False)
    fks = pd.read_csv(args.core_dir / "admin_fks.csv", dtype=str, keep_default_na=False)
    foreign_dependencies = pd.read_csv(args.core_dir / "foreign_dependencies.csv", dtype=str, keep_default_na=False)
    spatial = spatial_map(args.spatial_candidates)
    osm = read_osm_candidates(args.osm_admin)
    admin_rows, pcode_to_core = match_admin(sources, areas, names, spatial, osm)
    active_townships = areas[
        (areas["admin_level"] == "township")
        & areas["is_active"].map(bool_value)
        & (areas["deleted_at"] == "")
    ]
    baseline = pd.concat([
        baseline,
        pd.DataFrame([
            {
                "metric": "active_townships_marked_official_boundary",
                "value": str(int(active_townships["is_official_boundary"].map(bool_value).sum())),
                "detail": f"of {len(active_townships)} active townships",
            },
            {
                "metric": "active_townships_marked_public_usable",
                "value": str(int(active_townships["is_public_usable"].map(bool_value).sum())),
                "detail": f"of {len(active_townships)} active townships",
            },
            {
                "metric": "active_townships_marked_both_official_and_public",
                "value": str(int((active_townships["is_official_boundary"].map(bool_value) & active_townships["is_public_usable"].map(bool_value)).sum())),
                "detail": f"of {len(active_townships)} active townships",
            },
        ]),
    ], ignore_index=True)
    duplicates = duplicate_report(areas, names)
    en_by_code, my_by_code, postal_stats = read_postal(args.postal_zip)
    postal_rows, postal_manual = postal_audit(en_by_code, my_by_code, sources, towns, pcode_to_core)

    output = args.output_dir
    write_csv(output / "admin-match-report.csv", admin_rows)
    write_csv(output / "admin-manual-review.csv", [row for row in admin_rows if row["proposed_action"] in {"manual_review", "duplicate_merge_candidate", "blocked_missing_legal_geometry"}])
    write_csv(output / "admin-create-missing.csv", [row for row in admin_rows if row["proposed_action"] == "create_missing"])
    write_csv(output / "admin-duplicates.csv", duplicates)
    write_csv(output / "postal-match-report.csv", postal_rows)
    write_csv(output / "postal-manual-review.csv", postal_manual)
    (output / "summary.md").write_text(
        summary_markdown(admin_rows, postal_rows, postal_stats, baseline, foreign_dependencies, fks),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
